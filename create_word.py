import sys
import os

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
except ImportError:
    print("python-docx not installed. Exiting.")
    sys.exit(1)

doc = Document()

# Add Title
title = doc.add_heading('GenAI Lab: Live Presentation & Demo Script', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# Part 1
doc.add_heading('🎙️ Part 1: The Introduction (2-3 minutes)', level=1)
doc.add_heading('1. The Elevator Pitch:', level=2)
p = doc.add_paragraph()
p.add_run('"Welcome to the GenAI Lab! This project is an Autonomous AI Experimentation Engine. The goal of this platform is to completely automate the scientific method. Instead of humans spending weeks designing, running, and analyzing experiments, we use a council of specialized AI agents to generate hypotheses, debate experiment designs, run simulated trials in real-time, and write executive research reports."').italic = True

doc.add_heading('2. The Tech Stack:', level=2)
p = doc.add_paragraph()
p.add_run('"To achieve this, I built a modern, responsive web application. The frontend is built with React, Vite, and Framer Motion for premium, animated UI/UX, hosted on Vercel. The backend is powered by Python and FastAPI, hosted on Railway. For the artificial intelligence, I integrated both Groq\'s high-speed inference engine and Google\'s Gemini API to power the multi-agent system."').italic = True

# Part 2
doc.add_heading('💻 Part 2: The Live Testing Walkthrough (5-7 minutes)', level=1)
doc.add_paragraph("Now, share your screen and navigate through the live Vercel URL. Follow these exact steps:")

# Step 1
doc.add_heading('Step 1: The Dashboard (System Overview)', level=2)
doc.add_paragraph("Action: Start on the home page (/).", style='List Bullet')
doc.add_paragraph('Talking Point: "Here is the central command center. You can see our global UI design, which uses glassmorphism, responsive grids, and modern SVG icons. The dashboard tracks our success rates, active agents, and loops. Let\'s start a brand new research cycle."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Hover over the pulsing "Neural Engine Active" badge, point to the live date/time, and show the layout scale automatically when resizing the browser window.', style='List Bullet')

# Step 2
doc.add_heading('Step 2: The Hypothesis Lab (Idea Generation)', level=2)
doc.add_paragraph('Action: Navigate to Hypothesis Lab. Type a topic into the generator, for example: "The impact of quantum computing on machine learning algorithms." Click Generate.', style='List Bullet')
doc.add_paragraph('Talking Point: "First, we ask the AI to generate a scientifically sound hypothesis. It uses Groq to instantly formulate variables, a predicted outcome, and a risk level. I will review this and click Approve."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Read the generated Independent Variable (e.g. "Quantum Error Rates") and Dependent Variable (e.g. "Model Accuracy"), then click the green Approve button.', style='List Bullet')

# Step 3
doc.add_heading('Step 3: Experiment Center & Agent Debate', level=2)
doc.add_paragraph('Action: Navigate to Experiment Center. Find your approved hypothesis and click to design the experiment. Then, scroll down to the Agent Council / Debate section.', style='List Bullet')
doc.add_paragraph('Talking Point: "Now that we have a hypothesis, the system designs the testing parameters. But before we run it, our AI agents must debate it. Watch as different specialized agents (like the Data Scientist and the Ethics Reviewer) analyze the experiment from their unique perspectives."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Read one of the debate messages out loud (e.g., the Ethics agent pointing out potential computational resource biases) to show how autonomous agents interact.', style='List Bullet')

# Step 4
doc.add_heading('Step 4: Simulation Runner (Real-Time WebSockets)', level=2)
doc.add_paragraph('Action: Navigate to Simulation Runner. Select the experiment you just designed, choose a simulation type (e.g., Monte Carlo), set iterations to 1000, and click Run.', style='List Bullet')
doc.add_paragraph('Talking Point: "This is where the magic happens. When I start the simulation, the frontend connects to the FastAPI backend via WebSockets. The data streaming onto the screen—the live charts and progress bars—is happening in real-time. The server is crunching the numbers using Python\'s SciPy and NumPy libraries and pushing them instantly to the UI."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Point to the live line chart drawing itself dynamically as the progress bar fills up from 0 to 100%.', style='List Bullet')

# Step 5
doc.add_heading('Step 5: Analysis & Insights', level=2)
doc.add_paragraph('Action: Navigate to Analysis & Insights.', style='List Bullet')
doc.add_paragraph('Talking Point: "Once the simulation finishes, the AI analyzes the raw mathematical data. It calculates confidence scores and provides Explainable AI (XAI) insights so we know exactly why the experiment succeeded or failed."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Show the "Success" or "Failure" banner, and read one of the generated causal findings (e.g. "P-value < 0.05 confirming statistical significance").', style='List Bullet')

# Step 6
doc.add_heading('Step 6: Memory Vault (Knowledge Graph)', level=2)
doc.add_paragraph('Action: Navigate to Memory Vault. Search for the keyword "Quantum".', style='List Bullet')
doc.add_paragraph('Talking Point: "A true autonomous system must learn from its past. Everything we just did was automatically embedded and saved into the Memory Vault database. The system can reference these past successes and failures for future experiments."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Type "Quantum" into the search bar and watch the system instantly retrieve the exact experiment you just ran from SQLite.', style='List Bullet')

# Step 7
doc.add_heading('Step 7: Research Reports (The Grand Finale)', level=2)
doc.add_paragraph('Action: Navigate to Research Reports and click Generate AI Report.', style='List Bullet')
doc.add_paragraph('Talking Point: "Finally, we need to present our findings. The system aggregates the hypothesis, the agent debate, the simulation data, and the insights, and sends it all to Google\'s Gemini API. Gemini then writes a comprehensive, professional executive summary of the entire research cycle."', style='List Bullet')
doc.add_paragraph('Example Output to point out: Click generate, let the loading animation run, and scroll through the final generated markdown report showing formatting, tables, and bullet points.', style='List Bullet')

# Part 3
doc.add_heading('🎤 Part 3: The Conclusion & Q&A', level=1)
p = doc.add_paragraph()
p.add_run('"In summary, I have successfully deployed a full-stack, cloud-hosted application that orchestrates multiple LLMs to autonomously conduct scientific research. The frontend and backend are fully decoupled and communicate via REST APIs and real-time WebSockets.\n\nThank you for watching! I\'d be happy to answer any questions about the architecture, the deployment process, or the code."').italic = True

file_path = r'E:\4th Semester\Theory\Programming for AI\Project\GenAI Lab\GenAI_Lab_Presentation_Script.docx'
doc.save(file_path)
print("SUCCESS")
